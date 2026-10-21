"""Render the editable shapes and text in this presentation as a vector PDF."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "problem_solution.pptx"
OUTPUT = HERE / "problem_solution.pdf"
EMU_PER_POINT = 12700

FONT_DIR = Path("/System/Library/Fonts/Supplemental")
pdfmetrics.registerFont(TTFont("SlideArial", str(FONT_DIR / "Arial.ttf")))
pdfmetrics.registerFont(TTFont("SlideArialBold", str(FONT_DIR / "Arial Bold.ttf")))


def points(emu):
    return emu / EMU_PER_POINT


def rgb_color(color, fallback="#000000"):
    try:
        value = color.rgb
        if value is not None:
            return HexColor(f"#{value}")
    except (AttributeError, TypeError, ValueError):
        pass
    return HexColor(fallback)


def wrap(text, font, size, width):
    if not text:
        return [""]
    words = text.split()
    if not words:
        return [""]
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}" if current else word
        if current and pdfmetrics.stringWidth(candidate, font, size) > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    lines.append(current)
    return lines


def draw_text(pdf, shape, slide_height):
    frame = shape.text_frame
    if not frame.text:
        return
    left = points(shape.left + frame.margin_left)
    top = points(shape.top + frame.margin_top)
    box_width = points(shape.width - frame.margin_left - frame.margin_right)
    box_height = points(shape.height - frame.margin_top - frame.margin_bottom)
    prepared = []
    for paragraph in frame.paragraphs:
        font = "SlideArialBold" if paragraph.font.bold else "SlideArial"
        size = points(paragraph.font.size) if paragraph.font.size else 12
        color = rgb_color(paragraph.font.color, "#14243D")
        alignment = paragraph.alignment or PP_ALIGN.LEFT
        for line in wrap(paragraph.text, font, size, box_width):
            prepared.append((line, font, size, color, alignment, size * 1.12))
        if paragraph.space_after:
            prepared.append((None, font, size, color, alignment, points(paragraph.space_after)))
    total_height = sum(item[5] for item in prepared)
    if frame.vertical_anchor == MSO_ANCHOR.MIDDLE:
        cursor = slide_height - top - max(0, (box_height - total_height) / 2)
    elif frame.vertical_anchor == MSO_ANCHOR.BOTTOM:
        cursor = slide_height - top - max(0, box_height - total_height)
    else:
        cursor = slide_height - top
    for line, font, size, color, alignment, advance in prepared:
        if line is not None:
            pdf.setFillColor(color)
            pdf.setFont(font, size)
            line_width = pdfmetrics.stringWidth(line, font, size)
            draw_x = left
            if alignment == PP_ALIGN.CENTER:
                draw_x += (box_width - line_width) / 2
            elif alignment == PP_ALIGN.RIGHT:
                draw_x += box_width - line_width
            pdf.drawString(draw_x, cursor - size, line)
        cursor -= advance


def draw_shape(pdf, shape, slide_height):
    x = points(shape.left)
    y = slide_height - points(shape.top + shape.height)
    w = points(shape.width)
    h = points(shape.height)
    try:
        kind = shape.auto_shape_type
    except (AttributeError, ValueError):
        kind = None
    try:
        fill = rgb_color(shape.fill.fore_color) if shape.fill.type is not None else None
    except (AttributeError, TypeError, ValueError):
        fill = None
    try:
        stroke = rgb_color(shape.line.color) if shape.line.fill.type is not None else None
    except (AttributeError, TypeError, ValueError):
        stroke = None
    if fill or stroke:
        if fill:
            pdf.setFillColor(fill)
        if stroke:
            pdf.setStrokeColor(stroke)
            pdf.setLineWidth(points(shape.line.width) if shape.line.width else 0.8)
        use_fill = int(fill is not None)
        use_stroke = int(stroke is not None)
        if kind == MSO_SHAPE.OVAL:
            pdf.ellipse(x, y, x + w, y + h, stroke=use_stroke, fill=use_fill)
        elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            pdf.roundRect(x, y, w, h, min(w, h) * 0.12,
                          stroke=use_stroke, fill=use_fill)
        else:
            pdf.rect(x, y, w, h, stroke=use_stroke, fill=use_fill)
    if shape.has_text_frame:
        draw_text(pdf, shape, slide_height)


def main():
    deck = Presentation(SOURCE)
    width = points(deck.slide_width)
    height = points(deck.slide_height)
    pdf = canvas.Canvas(str(OUTPUT), pagesize=(width, height))
    pdf.setTitle("Медиаслед — проблематика и решение")
    pdf.setAuthor("Медиаслед")
    for slide in deck.slides:
        pdf.setFillColor(HexColor("#FFFFFF"))
        pdf.rect(0, 0, width, height, fill=1, stroke=0)
        for item in slide.shapes:
            draw_shape(pdf, item, height)
        pdf.showPage()
    pdf.save()
    print(OUTPUT)


if __name__ == "__main__":
    main()
