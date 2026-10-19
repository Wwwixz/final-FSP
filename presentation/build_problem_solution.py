"""Build the editable problem-and-solution section of the project presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "problem_solution.pptx"

NAVY = "0B1D36"
NAVY_2 = "152C4B"
INK = "14243D"
MUTED = "64738B"
RED = "E51D3B"
PINK = "FFF0F2"
BLUE = "EAF1FF"
PALE = "F4F7FB"
LINE = "E1E7EF"
WHITE = "FFFFFF"
GREEN = "218A62"


def rgb(value):
    return RGBColor.from_string(value)


def shape(slide, kind, x, y, w, h, fill, line=None, radius=False):
    item = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.line.fill.background() if line is None else None
    if line:
        item.line.color.rgb = rgb(line)
        item.line.width = Pt(1)
    if radius:
        item.adjustments[0] = 0.12
    return item


def txt(slide, value, x, y, w, h, size=18, color=INK, bold=False,
        align=PP_ALIGN.LEFT, font="Arial", margin=0, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    for idx, line in enumerate(value.split("\n")):
        para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        para.text = line
        para.alignment = align
        para.space_after = Pt(3)
        para.font.name = font
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = rgb(color)
    return box


def label(slide, value, x, y, w, fill=PINK, color=RED):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.32, fill, radius=True)
    txt(slide, value, x + 0.1, y, w - 0.2, 0.32, 10, color, True)


def page(slide, number, source="ТЗ: tz_final.pdf"):
    txt(slide, "МЕДИАСЛЕД", 0.68, 0.3, 2.7, 0.25, 10, MUTED, True)
    txt(slide, f"{source}  ·  {number:02d} / 06", 9.0, 7.13, 3.66, 0.2, 8, MUTED,
        align=PP_ALIGN.RIGHT)


def heading(slide, kicker, title, subtitle=None):
    label(slide, kicker, 0.68, 0.74, 1.64)
    txt(slide, title, 0.68, 1.15, 12.0, 0.68, 27, INK, True)
    if subtitle:
        txt(slide, subtitle, 0.7, 1.86, 11.8, 0.58, 13, MUTED)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# 01 — title
s = prs.slides.add_slide(blank)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, NAVY)
shape(s, MSO_SHAPE.RECTANGLE, 0.7, 1.28, 0.08, 3.75, RED)
label(s, "ПРОБЛЕМА И РЕШЕНИЕ", 0.96, 0.58, 2.31, "2C324A", WHITE)
txt(s, "От поиска постов —\nк проверяемой подборке", 0.96, 1.33, 8.0, 2.0, 35, WHITE, True)
txt(s, "Медиамониторинг VK, Одноклассников и MAX\nдля Агентства печати и массовых коммуникаций\nКрасноярского края", 0.98, 3.55, 7.5, 1.16, 17, "C7D3E2")
txt(s, "Фрагмент презентации · концепция продукта", 0.98, 6.69, 7, 0.3, 11, "8FA3BC")
for x, y, w, head, body in [
    (9.23, 1.38, 3.1, "ЗАПРОС", "Открытие ледового дворца"),
    (9.65, 2.86, 2.7, "ПОИСК", "VK · ОК · MAX"),
    (9.12, 4.34, 3.2, "РЕЗУЛЬТАТ", "Размещения + ссылки"),
]:
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 1.14, NAVY_2, "3D5270", True)
    txt(s, head, x + 0.18, y + 0.14, w - 0.36, 0.2, 9, "FB5E75", True)
    txt(s, body, x + 0.18, y + 0.39, w - 0.36, 0.53, 15, WHITE, True)

# 02 — what exactly breaks
s = prs.slides.add_slide(blank)
page(s, 2, "ТЗ, стр. 2 и 5")
heading(s, "ПРОБЛЕМА", "Одна новость — разные тексты и разные события",
        "Сотруднику нужно находить распространение конкретной новости, даже если её переписали.")
cards = [
    (0.68, "01", "Исходная публикация", "«В Красноярске открылся новый ледовый дворец»", BLUE, "4D73C8"),
    (4.89, "02", "Пересказ", "«Город получил новую ледовую арену»", PALE, NAVY_2),
    (9.10, "03", "Похожий текст", "«Ледовую арену открыли в другом городе»", PINK, RED),
]
for x, num, title, body, fill, accent in cards:
    shape(s, MSO_SHAPE.RECTANGLE, x, 2.76, 3.55, 2.47, fill, LINE, True)
    shape(s, MSO_SHAPE.OVAL, x + 0.22, 3.0, 0.45, 0.45, accent)
    txt(s, num, x + 0.22, 3.02, 0.45, 0.39, 12, WHITE, True, PP_ALIGN.CENTER)
    txt(s, title, x + 0.22, 3.58, 3.05, 0.39, 16, INK, True)
    txt(s, body, x + 0.22, 4.03, 3.05, 0.88, 15, MUTED)
txt(s, "✓", 4.32, 3.69, 0.42, 0.45, 22, GREEN, True, PP_ALIGN.CENTER)
txt(s, "×", 8.54, 3.69, 0.42, 0.45, 22, RED, True, PP_ALIGN.CENTER)
shape(s, MSO_SHAPE.RECTANGLE, 0.68, 5.67, 11.97, 0.94, NAVY, radius=True)
txt(s, "Простое совпадение слов не различает пересказ и другое событие.",
    1.0, 5.85, 11.3, 0.55, 19, WHITE, True)

# 03 — employee pain
s = prs.slides.add_slide(blank)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, NAVY)
label(s, "ЦЕНА РУЧНОЙ РАБОТЫ", 0.68, 0.72, 2.23, "2C324A", WHITE)
txt(s, "Сотрудник вынужден искать и перепроверять вручную", 0.68, 1.18, 11.9, 0.67, 27, WHITE, True)
txt(s, "Проблема возникает на каждом этапе — от охвата площадок до готового отчёта.",
    0.7, 1.91, 11.7, 0.45, 13, "B8C6D8")
pain = [
    (0.68, 2.68, "Разрозненные источники", "Публикации распределены между группами VK, ОК и каналами MAX; приходится повторять поиск."),
    (6.75, 2.68, "Пропуски и лишнее", "Перефразирование теряется, а общие слова приводят к материалам о другом событии."),
    (0.68, 4.48, "Ненадёжные сведения", "Нужно открыть ссылки и сверить дату, заголовок и выдержку с исходным постом."),
    (6.75, 4.48, "Неясный итог", "Повторный пост может удвоить счётчик, а сбой источника — выглядеть как отсутствие публикаций."),
]
for i, (x, y, title, body) in enumerate(pain):
    shape(s, MSO_SHAPE.RECTANGLE, x, y, 5.89, 1.52, NAVY_2, "39506E", True)
    shape(s, MSO_SHAPE.RECTANGLE, x + 0.18, y + 0.24, 0.06, 1.02, RED)
    txt(s, title, x + 0.38, y + 0.18, 5.17, 0.38, 17, WHITE, True)
    txt(s, body, x + 0.38, y + 0.61, 5.12, 0.75, 12, "C5D0DF")
txt(s, "Итог: подборка снова собирается вручную, хотя поиск уже был выполнен.",
    0.7, 6.55, 11.9, 0.43, 15, "F1A4B1", True)

# 04 — solution architecture
s = prs.slides.add_slide(blank)
page(s, 4, "Решение по ТЗ, стр. 3–6")
heading(s, "РЕШЕНИЕ", "Один запрос → проверяемые размещения",
        "Проектируемая система получает реальные посты, проверяет событие и сохраняет ссылки на оригиналы.")
stages = [
    (0.68, "1", "Запрос", "Тема, человек, событие\nили текст новости", BLUE),
    (3.82, "2", "Сбор", "Выбранные источники\nили широкий режим", PALE),
    (6.96, "3", "Отбор", "Смысл + факты события:\nместо, дата, участники", PALE),
    (10.10, "4", "Подборка", "Отдельные размещения\nи прямые ссылки", PINK),
]
for x, num, title, body, fill in stages:
    shape(s, MSO_SHAPE.RECTANGLE, x, 2.9, 2.55, 2.31, fill, LINE, True)
    shape(s, MSO_SHAPE.OVAL, x + 0.19, 3.1, 0.48, 0.48, RED if num == "4" else NAVY_2)
    txt(s, num, x + 0.19, 3.13, 0.48, 0.4, 12, WHITE, True, PP_ALIGN.CENTER)
    txt(s, title, x + 0.19, 3.72, 2.18, 0.4, 18, INK, True)
    txt(s, body, x + 0.19, 4.15, 2.18, 0.79, 12, MUTED)
for x in (3.36, 6.50, 9.64):
    txt(s, "→", x, 3.73, 0.43, 0.4, 24, RED, True, PP_ALIGN.CENTER)
shape(s, MSO_SHAPE.RECTANGLE, 0.68, 5.69, 11.97, 0.88, NAVY, radius=True)
txt(s, "Разные публикации считаем отдельно; повторное обнаружение того же поста не увеличивает итог.",
    0.98, 5.86, 11.35, 0.55, 16, WHITE, True)

# 05 — two modes and coverage
s = prs.slides.add_slide(blank)
page(s, 5, "ТЗ, стр. 4–6")
heading(s, "ОХВАТ", "Два режима поиска — с понятными границами",
        "Пользователь выбирает глубину поиска, а система показывает, что действительно обработано.")
for x, title, sub, bullets, fill in [
    (0.68, "По выбранным источникам", "Контролируемая область", [
        "Конкретные группы и каналы задаёт сотрудник",
        "У каждого источника виден результат обработки",
        "Подходит для регулярной проверки известного списка",
    ], BLUE),
    (6.75, "Шире списка", "По доступным публичным материалам", [
        "Ищем новые размещения вне справочника",
        "Период и Красноярский край сохраняются",
        "Ограничения площадок показываем явно",
    ], PALE),
]:
    shape(s, MSO_SHAPE.RECTANGLE, x, 2.63, 5.89, 3.10, fill, LINE, True)
    txt(s, title, x + 0.24, 2.91, 5.41, 0.44, 20, INK, True)
    txt(s, sub, x + 0.24, 3.39, 5.41, 0.28, 11, RED, True)
    for i, bullet in enumerate(bullets):
        shape(s, MSO_SHAPE.OVAL, x + 0.27, 3.97 + i * 0.46, 0.09, 0.09, RED)
        txt(s, bullet, x + 0.47, 3.86 + i * 0.46, 5.05, 0.32, 12, INK)
shape(s, MSO_SHAPE.RECTANGLE, 0.68, 6.02, 11.97, 0.53, PINK, radius=True)
txt(s, "Широкий режим расширяет охват, но не обещает просмотреть абсолютно все публикации соцсети.",
    0.96, 6.1, 11.4, 0.32, 12, RED, True)

# 06 — user-facing result and honest status
s = prs.slides.add_slide(blank)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, NAVY)
label(s, "РЕЗУЛЬТАТ ДЛЯ СОТРУДНИКА", 0.68, 0.7, 2.79, "2C324A", WHITE)
txt(s, "Подборка, которую можно быстро проверить", 0.68, 1.12, 11.96, 0.68, 27, WHITE, True)
for i, (title, body) in enumerate([
    ("Видно, что найдено", "Количество, источник, заголовок, дата или «недоступна», выдержка."),
    ("Видно, что не удалось", "Поиск, пустая выдача и ошибка источника имеют разные состояния."),
    ("Можно продолжить работу", "Найденное сохраняется при частичном сбое; XLSX повторяет подборку."),
]):
    y = 2.18 + i * 1.28
    shape(s, MSO_SHAPE.OVAL, 0.71, y + 0.12, 0.42, 0.42, RED)
    txt(s, str(i + 1), 0.71, y + 0.14, 0.42, 0.36, 11, WHITE, True, PP_ALIGN.CENTER)
    txt(s, title, 1.35, y, 5.38, 0.38, 17, WHITE, True)
    txt(s, body, 1.35, y + 0.41, 5.35, 0.61, 12, "C5D0DF")

shape(s, MSO_SHAPE.RECTANGLE, 7.02, 2.13, 5.55, 3.77, WHITE, radius=True)
txt(s, "Найдено 3 размещения", 7.30, 2.39, 4.98, 0.42, 18, INK, True)
for i, (headline, source) in enumerate([
    ("Открылся новый ледовый дворец", "VK · Источник 1"),
    ("Как выглядит новая ледовая арена", "ОК · Источник 2"),
    ("Первые тренировки на новой арене", "MAX · Источник 3"),
]):
    y = 2.97 + i * 0.87
    shape(s, MSO_SHAPE.RECTANGLE, 7.29, y, 5.0, 0.72, PALE, LINE, True)
    shape(s, MSO_SHAPE.RECTANGLE, 7.44, y + 0.14, 0.43, 0.43, "3D78D3" if i == 0 else "F38B32" if i == 1 else "6852CF", radius=True)
    txt(s, headline, 8.0, y + 0.08, 4.1, 0.32, 11, INK, True)
    txt(s, source, 8.0, y + 0.4, 4.1, 0.2, 9, MUTED)
txt(s, "Пример интерфейсной выдачи, не результаты реального поиска", 7.31, 5.55, 4.95, 0.22, 8, MUTED)
shape(s, MSO_SHAPE.RECTANGLE, 0.68, 6.48, 11.96, 0.47, "263A55", radius=True)
txt(s, "Текущий статус проекта: фронтенд и OpenAPI готовы; коннекторы соцсетей и оценка качества требуют реализации.",
    0.92, 6.56, 11.5, 0.31, 11, "DBE4EF")

prs.save(OUT)
print(OUT)
